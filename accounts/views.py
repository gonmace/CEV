import io
import logging
from functools import wraps

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model, login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.tokens import default_token_generator
from django.contrib.auth.views import LoginView, redirect_to_login
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.http import Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.encoding import force_bytes, force_str
from django.utils.http import (
    url_has_allowed_host_and_scheme, urlsafe_base64_decode, urlsafe_base64_encode,
)
from django.views.decorators.http import require_POST

from core.mail import send_mail_async

from django.db.models import Q
from django.db.models.functions import Lower

from .access import is_email_allowed, resolve_default_role, resolve_empresa
from .creditos import (
    anotar_creditos, fijar, modulos_activos, recargar, resumen_creditos,
    totales_en_circulacion, validar_tope,
)
from .marca import (
    aplicar_colores, con_logo, insertar_logo, logo_de, marca_de, marca_propia,
    placeholders_de, plantilla_de, puede_editar_marca, reemplazar_placeholders,
)
from .forms import (
    ActivationForm, AdminUserEditForm, AllowedEmailForm, EmailAuthenticationForm,
    EmailConfigForm, EmpresaForm, InviteForm, MarcaForm, PlantillasForm,
    ProfileNameForm, RequestAccessForm, role_choices_for,
)
from .models import (
    AllowedEmail, EmailConfig, Empresa, Profile, Role, RolePermission, TopeUsuario,
    UserPermission,
)
from .permissions import (
    CAPABILITIES, MODULE_KEYS, MODULES, anotar_modulos, get_user_empresa, get_user_role,
    has_capability, is_company_admin, puede_acceder_a,
)

User = get_user_model()

# La impersonación se audita por acá (ver views.impersonate): quién actuó como quién.
logger = logging.getLogger(__name__)

NEUTRAL_MSG = (
    'Si el correo está habilitado, te enviamos un enlace para activar tu cuenta. '
    'Revisa tu bandeja de entrada.'
)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _get_or_create_pending_user(email, role=''):
    """Obtiene/crea un usuario inactivo para el correo. Si se pasa rol, fija el Profile."""
    user = User.objects.filter(email__iexact=email).first()
    if user is None:
        user = User.objects.create(username=email[:150], email=email, is_active=False)
        user.set_unusable_password()
        user.save()
    if role:
        Profile.objects.update_or_create(user=user, defaults={'role': role})
    return user


def _activation_path(user):
    uid = urlsafe_base64_encode(force_bytes(user.pk))
    token = default_token_generator.make_token(user)
    return reverse('accounts:activate', args=[uid, token])


def _send_activation_email(request, user):
    link = request.build_absolute_uri(_activation_path(user))
    body = render_to_string('accounts/emails/activation.txt', {'user': user, 'link': link})
    send_mail_async('Activa tu cuenta de CEV', body, [user.email])


def _superuser_required(view):
    # No usar user_passes_test tal cual: si el usuario ya está autenticado pero no es
    # superusuario, redirigir al login (con CustomLoginView.redirect_authenticated_user=True)
    # genera un ping-pong infinito de redirects login↔admin. Un no-superusuario autenticado
    # debe recibir 403; solo el anónimo va al login.
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if request.user.is_superuser:
            return view(request, *args, **kwargs)
        if request.user.is_authenticated:
            raise PermissionDenied
        return redirect_to_login(request.get_full_path(), reverse('accounts:login'))
    return wrapped


def _company_admin_required(view):
    # Mismo criterio que _superuser_required (403 al autenticado sin permiso, login al
    # anónimo), pero para el Administrador de una empresa en vez del superuser.
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if is_company_admin(request.user):
            return view(request, *args, **kwargs)
        if request.user.is_authenticated:
            raise PermissionDenied
        return redirect_to_login(request.get_full_path(), reverse('accounts:login'))
    return wrapped


def _client_ip(request):
    # X-Real-IP (no XFF): el primer hop de X-Forwarded-For lo controla el cliente —
    # con XFF el throttle por IP se bypaseaba mandando un header falso por request.
    return request.META.get('HTTP_X_REAL_IP') or request.META.get('REMOTE_ADDR', '')


def _request_access_allowed(request, email):
    """Throttle anti-abuso: cooldown por email (10 min) + tope por IP (10/h).

    Usa django.core.cache — sin CACHES definido cae a locmem (por proceso), que
    alcanza como mitigación."""
    ip = _client_ip(request)
    email_key = f'reqacc:email:{email.lower()}'
    ip_key = f'reqacc:ip:{ip}'
    if cache.get(email_key):
        return False
    ip_count = cache.get(ip_key, 0)
    if ip_count >= 10:
        return False
    cache.set(email_key, 1, 600)            # 10 minutos
    cache.set(ip_key, ip_count + 1, 3600)   # ventana de 1 hora
    return True


# ── Onboarding ──────────────────────────────────────────────────────────────

def request_access(request):
    if request.user.is_authenticated:
        return redirect('home')
    if request.method == 'POST':
        form = RequestAccessForm(request.POST)
        if form.is_valid():
            email = form.cleaned_data['email']
            # Throttle + allow-list. Solo se envía activación a cuentas que NUNCA activaron
            # (sin contraseña usable): así un usuario dado de baja no puede reactivarse solo.
            if _request_access_allowed(request, email) and is_email_allowed(email):
                user = _get_or_create_pending_user(email)
                if not user.is_active and not user.has_usable_password():
                    if resolve_empresa(email):
                        # Correo de dominio de empresa: el dominio ya lo certifica, así que
                        # se salta la verificación por correo y se va directo a fijar
                        # password (decisión de producto — a costa de la protección
                        # anti-enumeración que sí aplica al resto de este flujo).
                        return redirect(_activation_path(user))
                    _send_activation_email(request, user)
            messages.success(request, NEUTRAL_MSG)
            return redirect('accounts:login')
    else:
        form = RequestAccessForm()
    return render(request, 'accounts/request_access.html', {'form': form})


def activate(request, uidb64, token):
    try:
        uid = force_str(urlsafe_base64_decode(uidb64))
        user = User.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        user = None

    valid = user is not None and default_token_generator.check_token(user, token)
    # Defensa en profundidad: una cuenta inactiva pero CON contraseña usable fue dada de baja
    # por el admin; no debe poder auto-reactivarse por este flujo (solo el admin la reactiva).
    if user is not None and not user.is_active and user.has_usable_password():
        valid = False
    if not valid:
        return render(request, 'accounts/activate.html', {'invalid': True})

    if request.method == 'POST':
        form = ActivationForm(user, request.POST)
        if form.is_valid():
            # Última comprobación de cupo: entre la invitación y este clic pueden haberse
            # ocupado las plazas de la empresa.
            empresa = resolve_empresa(user.email)
            if empresa and not empresa.hay_cupo:
                messages.error(
                    request,
                    f'«{empresa.nombre}» alcanzó su límite de usuarios. '
                    'Pide al administrador que amplíe el cupo.',
                )
                return render(request, 'accounts/activate.html', {
                    'form': form, 'invalid': False, 'email': user.email,
                })
            form.save()
            user.is_active = True
            user.save(update_fields=['is_active'])
            # La empresa se resuelve por el dominio del correo: sin empresa, la cuenta es
            # personal (créditos propios, sin equipo). El rol solo se fija si el Profile no
            # existía: si vino de una invitación, ya trae el rol que eligió el superuser y
            # no hay que pisarlo.
            profile, _ = Profile.objects.get_or_create(
                user=user, defaults={'role': resolve_default_role(user.email)},
            )
            if profile.empresa_id != (empresa.id if empresa else None):
                profile.empresa = empresa
                profile.save(update_fields=['empresa'])
            login(request, user, backend='accounts.backends.EmailBackend')
            messages.success(request, '¡Cuenta activada! Bienvenido/a.')
            return redirect('home')
    else:
        form = ActivationForm(user)
    return render(request, 'accounts/activate.html', {'form': form, 'invalid': False, 'email': user.email})


@login_required
def profile(request):
    if request.method == 'POST':
        form = ProfileNameForm(request.POST, instance=request.user)
        if form.is_valid():
            form.save()
            messages.success(request, 'Perfil actualizado.')
            return redirect('accounts:profile')
    else:
        form = ProfileNameForm(instance=request.user)
    return render(request, 'accounts/profile.html', {'form': form})


class CustomLoginView(LoginView):
    template_name = 'accounts/login.html'
    authentication_form = EmailAuthenticationForm
    redirect_authenticated_user = True

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        # El botón de Google solo se muestra si la SocialApp está cargada en la BD —
        # sin credenciales, allauth responde 500 al iniciar el flujo OAuth.
        from allauth.socialaccount.models import SocialApp
        ctx['google_login_enabled'] = SocialApp.objects.filter(provider='google').exists()
        return ctx

    def form_valid(self, form):
        response = super().form_valid(form)
        if not form.cleaned_data.get('remember_me'):
            self.request.session.set_expiry(0)  # expira al cerrar el navegador
        return response


# ── Gestión de cuentas ────────────────────────────────────────────────────────

@_superuser_required
def access_admin(request):
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'add_empresa':
            form = EmpresaForm(request.POST)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.created_by = request.user
                obj.save()
                messages.success(request, f'Empresa «{obj.nombre}» habilitada.')
            else:
                messages.error(request, 'Revisa los datos de la empresa.')
        elif action == 'add_email':
            form = AllowedEmailForm(request.POST, viewer=request.user)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.created_by = request.user
                obj.save()
                messages.success(request, f'Correo «{obj.email}» agregado.')
            else:
                messages.error(request, 'Revisa los datos del correo.')
        elif action == 'invite':
            form = InviteForm(request.POST, viewer=request.user)
            if form.is_valid():
                email = form.cleaned_data['email']
                role = form.cleaned_data['role']
                # Si el correo es de una empresa habilitada, la invitación ocupa una de sus
                # plazas: no se puede invitar por encima del cupo contratado.
                empresa = resolve_empresa(email)
                if empresa and not empresa.hay_cupo:
                    messages.error(
                        request,
                        f'«{empresa.nombre}» alcanzó su límite de {empresa.max_usuarios} '
                        'usuarios. Amplía el cupo o desactiva una cuenta antes de invitar.',
                    )
                else:
                    AllowedEmail.objects.update_or_create(
                        email=email,
                        defaults={'default_role': role, 'is_active': True, 'created_by': request.user},
                    )
                    user = _get_or_create_pending_user(email, role=role)
                    if user.is_active:
                        messages.info(request, f'«{email}» ya tiene cuenta activa.')
                    else:
                        _send_activation_email(request, user)
                        destino = f' ({empresa.nombre})' if empresa else ''
                        messages.success(request, f'Invitación enviada a «{email}»{destino}.')
            else:
                messages.error(request, 'Correo inválido para invitar.')
        elif action == 'editar_empresa':
            # Editor de la empresa (modal): todos sus datos y el saldo de cada módulo de una
            # sola vez. Los créditos se FIJAN al valor escrito (no se suman): se escribe el
            # saldo que debe quedar, que es como uno piensa el plan contratado.
            empresa = get_object_or_404(Empresa, pk=request.POST.get('id'))
            nombre = (request.POST.get('nombre') or empresa.nombre).strip()
            dominio = (request.POST.get('dominio') or empresa.dominio).strip().lower().lstrip('@')

            try:
                max_usuarios = int(request.POST.get('max_usuarios') or empresa.max_usuarios)
            except ValueError:
                max_usuarios = empresa.max_usuarios

            # El dominio es la clave por la que se resuelve a qué empresa entra cada correo,
            # así que no puede chocar con el de otra.
            dominio_ocupado = (
                Empresa.objects.filter(dominio=dominio).exclude(pk=empresa.pk).exists()
            )

            if max_usuarios < empresa.usuarios_activos:
                messages.error(
                    request,
                    f'«{empresa.nombre}» ya tiene {empresa.usuarios_activos} cuentas activas: '
                    f'no se puede bajar el cupo a {max_usuarios}. Desactiva cuentas primero.',
                )
            elif dominio_ocupado:
                messages.error(request, f'El dominio «{dominio}» ya lo usa otra empresa.')
            else:
                empresa.nombre = nombre
                empresa.dominio = dominio
                empresa.max_usuarios = max_usuarios
                empresa.save(update_fields=['nombre', 'dominio', 'max_usuarios'])
                for modulo, _label in MODULES:
                    valor = request.POST.get(f'creditos_{modulo}')
                    if valor not in (None, ''):
                        try:
                            fijar(empresa, modulo, int(valor))
                        except ValueError:
                            pass
                messages.success(request, f'«{empresa.nombre}» actualizada.')
        elif action == 'toggle_empresa':
            obj = get_object_or_404(Empresa, pk=request.POST.get('id'))
            obj.is_active = not obj.is_active
            obj.save(update_fields=['is_active'])
        elif action == 'toggle_email':
            obj = get_object_or_404(AllowedEmail, pk=request.POST.get('id'))
            obj.is_active = not obj.is_active
            obj.save(update_fields=['is_active'])
        elif action == 'delete_empresa':
            empresa = get_object_or_404(Empresa, pk=request.POST.get('id'))
            if empresa.miembros.exists():
                messages.error(
                    request,
                    f'«{empresa.nombre}» tiene {empresa.miembros.count()} cuenta(s) asociada(s). '
                    'Desactivala en vez de eliminarla para no dejarlas huérfanas.',
                )
            else:
                nombre = empresa.nombre
                empresa.delete()
                messages.success(request, f'Empresa «{nombre}» eliminada.')
        elif action == 'recargar_creditos':
            # Recarga manual: el superuser suma créditos a la bolsa de una empresa o de
            # una cuenta personal. No hay pasarela de pago; el alta es comercial.
            modulo = request.POST.get('modulo', '')
            try:
                cantidad = int(request.POST.get('cantidad', 0))
            except (TypeError, ValueError):
                cantidad = 0
            if modulo not in MODULE_KEYS:
                messages.error(request, 'Módulo inválido.')
            elif cantidad <= 0:
                messages.error(request, 'La cantidad de créditos debe ser mayor que cero.')
            else:
                empresa_id = request.POST.get('empresa_id')
                usuario_id = request.POST.get('usuario_id')
                destino = (
                    get_object_or_404(Empresa, pk=empresa_id) if empresa_id
                    else get_object_or_404(User, pk=usuario_id)
                )
                bolsa = recargar(destino, modulo, cantidad)
                etiqueta = dict(MODULES)[modulo]
                messages.success(
                    request,
                    f'+{cantidad} créditos de {etiqueta} para «{destino}». '
                    f'Saldo: {bolsa.creditos}.',
                )
        elif action == 'delete_email':
            get_object_or_404(AllowedEmail, pk=request.POST.get('id')).delete()
            messages.success(request, 'Correo eliminado.')
        elif action == 'set_email_role':
            obj = get_object_or_404(AllowedEmail, pk=request.POST.get('id'))
            role = request.POST.get('default_role', '')
            # role_choices_for: para un no-superuser, ADMINISTRADOR no es un rol válido.
            if role and role not in dict(role_choices_for(request.user)):
                messages.error(request, 'Rol inválido.')
            else:
                obj.default_role = role  # '' = sin rol (cae al default del dominio / USUARIO)
                obj.save(update_fields=['default_role'])
                messages.success(request, f'Rol de «{obj.email}» actualizado.')
        elif action == 'toggle_user':
            target = get_object_or_404(User, pk=request.POST.get('id'))
            if target.pk == request.user.pk or target.is_superuser:
                messages.error(request, 'No puedes cambiar el estado de este usuario.')
            elif not target.is_active and not target.has_usable_password():
                # Todavía no activó su cuenta (sin contraseña): no se puede activar a mano,
                # tiene que pasar por el enlace de invitación (ver activate()/request_access()).
                messages.error(
                    request,
                    f'«{target.email or target.username}» todavía no activó su cuenta. '
                    'Usa «Reenviar invitación» en vez de activarla directamente.',
                )
            else:
                target.is_active = not target.is_active
                target.save(update_fields=['is_active'])
                estado = 'activado' if target.is_active else 'desactivado'
                messages.success(request, f'Usuario «{target.email or target.username}» {estado}.')
        elif action == 'toggle_module':
            # Toggle de acceso a un módulo desde la tabla de cuentas. Escribe un override
            # individual (UserPermission) con el valor contrario al efectivo actual: no toca
            # el default del rol, que se sigue editando en el tablero de Roles.
            target = get_object_or_404(User, pk=request.POST.get('id'))
            capability = request.POST.get('capability', '')
            if capability not in MODULE_KEYS:
                messages.error(request, 'Módulo inválido.')
            elif target.is_superuser:
                # El superuser saltea todos los checks: un override sobre él no haría nada.
                messages.error(request, 'El superuser tiene acceso a todos los módulos.')
            else:
                nuevo = not has_capability(target, capability)
                UserPermission.objects.update_or_create(
                    user=target, capability=capability, defaults={'enabled': nuevo},
                )
                etiqueta = dict(MODULES)[capability]
                estado = 'habilitado' if nuevo else 'deshabilitado'
                messages.success(
                    request,
                    f'Módulo {etiqueta} {estado} para «{target.email or target.username}».',
                )
        elif action == 'set_tope':
            # Tope individual editado directamente desde "Cuentas registradas" — evita el
            # viaje a la ficha de la cuenta solo para cambiar un número. El superuser no
            # tiene tope: su bypass es incondicional (ver accounts.creditos._es_ilimitado),
            # así que acá se lo rechaza en vez de guardar un valor que nunca aplicaría.
            target = get_object_or_404(User, pk=request.POST.get('id'))
            modulo = request.POST.get('modulo', '')
            if modulo not in MODULE_KEYS:
                messages.error(request, 'Módulo inválido.')
            elif target.is_superuser:
                messages.error(request, 'El superuser no consume créditos: no tiene tope.')
            else:
                tope_raw = (request.POST.get('tope') or '').strip()
                if tope_raw:
                    try:
                        tope = max(0, int(tope_raw))
                    except ValueError:
                        messages.error(request, 'El tope debe ser un número.')
                    else:
                        error = validar_tope(target, modulo, tope)
                        if error:
                            messages.error(request, error)
                        else:
                            TopeUsuario.objects.update_or_create(
                                usuario=target, modulo=modulo,
                                defaults={'tope': tope},
                            )
                            messages.success(
                                request,
                                f'Tope de «{target.email or target.username}» actualizado a {tope}.',
                            )
                else:
                    TopeUsuario.objects.filter(usuario=target, modulo=modulo).delete()
                    messages.success(
                        request,
                        f'Tope de «{target.email or target.username}» quitado (usa toda la bolsa compartida).',
                    )
        elif action == 'resend_invite':
            target = get_object_or_404(User, pk=request.POST.get('id'))
            if target.is_active:
                messages.info(request, 'Ese usuario ya tiene la cuenta activa.')
            else:
                _send_activation_email(request, target)
                messages.success(request, f'Invitación reenviada a «{target.email}».')
        elif action == 'delete_user':
            target = get_object_or_404(User, pk=request.POST.get('id'))
            label = target.email or target.username
            if target.pk == request.user.pk or target.is_superuser:
                messages.error(request, 'No puedes eliminar este usuario.')
            else:
                from pliego_licitacion.models import EspecificacionTecnica
                from proyectos.models import Proyecto
                from servicios.models import Servicio
                tiene_contenido = (
                    Proyecto.objects.filter(creado_por=target).exists()
                    or Servicio.objects.filter(creado_por=target).exists()
                    or EspecificacionTecnica.objects.filter(creado_por=target).exists()
                )
                if tiene_contenido:
                    messages.error(
                        request,
                        f'«{label}» tiene proyectos, servicios o especificaciones creados. '
                        'Desactiva la cuenta en vez de eliminarla para no perder la autoría.',
                    )
                else:
                    target.delete()
                    messages.success(request, f'Cuenta «{label}» eliminada.')
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            pending = list(messages.get_messages(request))
            last = pending[-1] if pending else None
            return JsonResponse({'message': str(last) if last else '', 'tag': last.tags if last else ''})
        return redirect('accounts:access_admin')

    users_qs = (
        User.objects.select_related('profile', 'profile__empresa')
        .annotate(empresa_nombre_lower=Lower('profile__empresa__nombre'))
        .order_by('empresa_nombre_lower', 'email')
    )
    empresas = Empresa.objects.prefetch_related('bolsas').all()
    emails = AllowedEmail.objects.all()
    users_page = Paginator(users_qs, 20).get_page(request.GET.get('page'))
    # Estado efectivo de módulos y créditos de la página (unas pocas queries, no un N+1
    # por celda). Ver permissions.anotar_modulos y creditos.anotar_creditos.
    anotar_modulos(users_page.object_list)
    anotar_creditos(users_page.object_list)

    # Saldo por módulo de cada empresa, para pintarlo junto a su fila sin un N+1 por celda.
    for empresa in empresas:
        saldos = {b.modulo: b.creditos for b in empresa.bolsas.all()}
        empresa.creditos_por_modulo = [
            {'key': key, 'label': label, 'saldo': saldos.get(key, 0)}
            for key, label in MODULES
        ]

    # Barra de resumen: los números que el superuser quiere ver de un vistazo al entrar.
    # «Pendientes» son las cuentas invitadas que todavía no activaron (sin contraseña
    # usable): son las que pueden necesitar un reenvío de invitación.
    circulacion = totales_en_circulacion()
    resumen = {
        'empresas_activas': Empresa.objects.filter(is_active=True).count(),
        'cuentas_activas': User.objects.filter(is_active=True).count(),
        # Invitadas pero sin activar: inactivas y sin contraseña usable (Django marca esas
        # con '!' delante). Es el mismo criterio con el que activate() decide si una cuenta
        # nunca llegó a activarse o fue dada de baja por el admin.
        'pendientes': User.objects.filter(is_active=False).filter(
            Q(password='') | Q(password__startswith='!')
        ).count(),
        'creditos': [
            {'key': key, 'label': label, 'total': circulacion.get(key, 0)}
            for key, label in MODULES
        ],
    }

    return render(request, 'accounts/access_admin.html', {
        'empresas': empresas,
        'emails': emails,
        'modulos': MODULES,
        'resumen': resumen,
        'role_choices': role_choices_for(request.user),
        'users': users_page,
        'empresa_form': EmpresaForm(),
        'email_form': AllowedEmailForm(viewer=request.user),
        'invite_form': InviteForm(viewer=request.user),
    })


@_company_admin_required
def mi_equipo(request):
    """Panel acotado para que un Administrador de empresa gestione SU equipo: invitar
    gente (siempre del dominio de su empresa), promover/degradar el rol de sus
    compañeros y repartir el tope de créditos de cada uno sobre la bolsa compartida.

    No puede recargar la bolsa (eso es comercial, sigue siendo del superuser en Cuentas)
    ni tocar cuentas de otra empresa — todo acá se valida contra `empresa`, la suya."""
    empresa = get_user_empresa(request.user)

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'invite':
            form = InviteForm(request.POST, full_roles=True)
            if form.is_valid():
                email = form.cleaned_data['email']
                role = form.cleaned_data['role']
                _, _, domain = email.rpartition('@')
                if domain != empresa.dominio:
                    messages.error(
                        request,
                        f'Solo puedes invitar correos del dominio @{empresa.dominio}.',
                    )
                elif not empresa.hay_cupo:
                    messages.error(
                        request,
                        f'«{empresa.nombre}» alcanzó su límite de {empresa.max_usuarios} '
                        'usuarios. Pídele al administrador que amplíe el cupo.',
                    )
                else:
                    AllowedEmail.objects.update_or_create(
                        email=email,
                        defaults={'default_role': role, 'is_active': True, 'created_by': request.user},
                    )
                    user = _get_or_create_pending_user(email, role=role)
                    if user.is_active:
                        messages.info(request, f'«{email}» ya tiene cuenta activa.')
                    else:
                        _send_activation_email(request, user)
                        messages.success(request, f'Invitación enviada a «{email}».')
            else:
                messages.error(request, 'Correo inválido para invitar.')
        elif action == 'toggle_role':
            # Solo sobre miembros de la MISMA empresa: sin este chequeo, un Administrador
            # podría mandar el id de cualquier usuario del sistema.
            target = get_object_or_404(User, pk=request.POST.get('id'))
            if target.pk == request.user.pk or target.is_superuser or get_user_empresa(target) != empresa:
                messages.error(request, 'No puedes cambiar el rol de esta cuenta.')
            else:
                nuevo_rol = Role.USUARIO if get_user_role(target) == Role.ADMINISTRADOR else Role.ADMINISTRADOR
                Profile.objects.filter(user=target).update(role=nuevo_rol)
                messages.success(
                    request,
                    f'«{target.email or target.username}» ahora es {dict(Role.choices)[nuevo_rol]}.',
                )
        elif action == 'set_tope':
            target = get_object_or_404(User, pk=request.POST.get('id'))
            modulo = request.POST.get('modulo', '')
            if get_user_empresa(target) != empresa:
                messages.error(request, 'Esa cuenta no pertenece a tu empresa.')
            elif modulo not in MODULE_KEYS:
                messages.error(request, 'Módulo inválido.')
            else:
                tope_raw = (request.POST.get('tope') or '').strip()
                if tope_raw:
                    try:
                        tope = max(0, int(tope_raw))
                    except ValueError:
                        messages.error(request, 'El tope debe ser un número.')
                    else:
                        error = validar_tope(target, modulo, tope)
                        if error:
                            messages.error(request, error)
                        else:
                            TopeUsuario.objects.update_or_create(
                                usuario=target, modulo=modulo,
                                defaults={'tope': tope},
                            )
                            messages.success(
                                request,
                                f'Tope de «{target.email or target.username}» actualizado a {tope}.',
                            )
                else:
                    TopeUsuario.objects.filter(usuario=target, modulo=modulo).delete()
                    messages.success(
                        request,
                        f'Tope de «{target.email or target.username}» quitado (usa toda la bolsa compartida).',
                    )
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            pending = list(messages.get_messages(request))
            last = pending[-1] if pending else None
            return JsonResponse({'message': str(last) if last else '', 'tag': last.tags if last else ''})
        return redirect('accounts:mi_equipo')

    miembros = anotar_creditos(list(
        User.objects.filter(profile__empresa=empresa)
        .select_related('profile')
        .order_by('is_active', 'email')
    ))
    saldos = {b.modulo: b.creditos for b in empresa.bolsas.all()}
    creditos_empresa = [
        {'key': key, 'label': label, 'saldo': saldos.get(key, 0)}
        for key, label in MODULES
    ]

    return render(request, 'accounts/mi_equipo.html', {
        'empresa': empresa,
        'miembros': miembros,
        'creditos_empresa': creditos_empresa,
        'modulos': MODULES,
        'invite_form': InviteForm(full_roles=True),
    })


@_superuser_required
def user_edit(request, pk):
    target = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action', 'save_account')
        # Cualquier cuenta admite overrides individuales, sea cual sea su rol. Editarlos
        # es gestión de privilegios: queda reservado al superuser (el gestor con
        # accounts.manage solo edita nombre/rol). Se revalida acá server-side, no solo
        # ocultando la sección en el template.
        if action in ('save_permissions', 'reset_permissions'):
            if not request.user.is_superuser:
                raise PermissionDenied
            if action == 'reset_permissions':
                UserPermission.objects.filter(user=target).delete()
                messages.success(request, 'Se restableció el default del rol.')
            else:
                for cap_key, _label in CAPABILITIES:
                    enabled = request.POST.get(cap_key) == 'on'
                    UserPermission.objects.update_or_create(
                        user=target, capability=cap_key, defaults={'enabled': enabled},
                    )
                messages.success(request, 'Permisos personalizados de la cuenta actualizados.')
            return redirect('accounts:user_edit', pk=target.pk)

        if action == 'guardar_creditos':
            # Dos cosas distintas según el tipo de cuenta:
            #  - personal: se recarga SU bolsa (es la única que tiene).
            #  - de empresa: no se recarga nada acá (la bolsa es de la empresa, se recarga
            #    en Cuentas); lo que se fija es su TOPE de gasto sobre esa bolsa común.
            empresa_target = getattr(getattr(target, 'profile', None), 'empresa', None)
            for modulo, etiqueta in MODULES:
                if empresa_target is None:
                    try:
                        recarga = int(request.POST.get(f'recarga_{modulo}') or 0)
                    except ValueError:
                        recarga = 0
                    if recarga > 0:
                        recargar(target, modulo, recarga)

                tope_raw = (request.POST.get(f'tope_{modulo}') or '').strip()
                if tope_raw:
                    try:
                        tope = max(0, int(tope_raw))
                    except ValueError:
                        continue
                    # Se valida DESPUÉS de la recarga: recargar 100 y fijar tope 100 en
                    # el mismo guardado es legítimo.
                    error = validar_tope(target, modulo, tope)
                    if error:
                        messages.error(request, f'{etiqueta}: {error}')
                        continue
                    TopeUsuario.objects.update_or_create(
                        usuario=target, modulo=modulo,
                        defaults={'tope': tope},
                    )
                else:
                    # Campo vacío = sin tope: puede gastar toda la bolsa.
                    TopeUsuario.objects.filter(usuario=target, modulo=modulo).delete()
            messages.success(request, 'Créditos actualizados.')
            return redirect('accounts:user_edit', pk=target.pk)

        form = AdminUserEditForm(request.POST, instance=target, viewer=request.user)
        if form.is_valid():
            new_role = form.cleaned_data['role']
            # El superuser es un rol de plataforma, no de una empresa: puede_ver_todo() ya
            # le da acceso a los documentos de todas por igual, así que asignarle una lo
            # metería en un compartimento del que no forma parte.
            new_empresa = form.cleaned_data['empresa'] if not target.is_superuser else None
            profile = getattr(target, 'profile', None)
            empresa_actual_id = profile.empresa_id if profile else None
            # Cupo: solo se revisa si de verdad cambia de empresa (asignarla a la que ya
            # tiene, o quitársela, nunca debe bloquearse por esto).
            if new_empresa and new_empresa.id != empresa_actual_id and not new_empresa.hay_cupo:
                messages.error(
                    request,
                    f'«{new_empresa.nombre}» alcanzó su límite de {new_empresa.max_usuarios} '
                    'usuarios. Amplía el cupo o desactiva una cuenta antes de asignar esta.',
                )
                return redirect('accounts:user_edit', pk=target.pk)
            form.save()
            Profile.objects.update_or_create(
                user=target, defaults={'role': new_role, 'empresa': new_empresa},
            )
            # Los overrides sobreviven al cambio de rol: son decisiones sobre esta cuenta,
            # no sobre su rol. Cambiar el rol cambia el default; lo que estaba pisado sigue
            # pisado (se ve en la ficha, y «Restablecer default» los limpia).
            messages.success(request, f'Cuenta «{target.email or target.username}» actualizada.')
            return redirect('accounts:access_admin')
    else:
        form = AdminUserEditForm(instance=target, viewer=request.user)

    target_role = get_user_role(target)
    # Sobre un superuser no se muestran: saltea todos los checks (has_capability), así que
    # cualquier toggle sería una configuración fantasma que no cambia nada.
    show_overrides = request.user.is_superuser and not target.is_superuser
    permission_rows = []
    if show_overrides:
        overrides = {up.capability: up.enabled for up in UserPermission.objects.filter(user=target)}
        role_defaults = {
            rp.capability: rp.enabled for rp in RolePermission.objects.filter(role=target_role)
        }
        for cap_key, cap_label in CAPABILITIES:
            permission_rows.append({
                'key': cap_key, 'label': cap_label,
                'enabled': overrides.get(cap_key, role_defaults.get(cap_key, False)),
                'role_default': role_defaults.get(cap_key, False),
                'is_overridden': cap_key in overrides,
            })

    empresa_target = getattr(getattr(target, 'profile', None), 'empresa', None)

    return render(request, 'accounts/user_edit.html', {
        'form': form, 'target': target,
        'show_overrides': show_overrides,
        'permission_rows': permission_rows,
        'has_overrides': any(r['is_overridden'] for r in permission_rows),
        'empresa_target': empresa_target,
        # En una cuenta de empresa el saldo que se ve es el de la bolsa compartida; en una
        # personal, el suyo propio. `resumen_creditos` ya resuelve cuál corresponde.
        'creditos_rows': resumen_creditos(target) if not target.is_superuser else [],
    })


@_superuser_required
def roles_board(request):
    roles = Role.choices  # [(value, label), ...]

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'save_matrix':
            for role_value, _ in roles:
                for cap_key, _label in CAPABILITIES:
                    enabled = request.POST.get(f'{role_value}:{cap_key}') == 'on'
                    RolePermission.objects.update_or_create(
                        role=role_value, capability=cap_key, defaults={'enabled': enabled},
                    )
            messages.success(request, 'Permisos actualizados.')
        return redirect('accounts:roles_board')

    # Matriz actual {(role, cap): enabled}
    current = {
        (rp.role, rp.capability): rp.enabled
        for rp in RolePermission.objects.all()
    }
    matrix = []
    for cap_key, cap_label in CAPABILITIES:
        row = {'key': cap_key, 'label': cap_label, 'cells': []}
        for role_value, role_label in roles:
            row['cells'].append({
                'role': role_value,
                'label': role_label,
                'enabled': current.get((role_value, cap_key), False),
            })
        matrix.append(row)

    return render(request, 'accounts/roles_board.html', {
        'roles': roles,
        'matrix': matrix,
    })


def _send_test_email(request, data):
    """Envía un correo de prueba al superuser logueado con la config del form (sin
    guardar): SMTP propio si `enabled` + host, o el backend del settings si no —
    exactamente lo que quedaría efectivo al guardar."""
    from django.core.mail import get_connection, send_mail
    to = request.user.email
    if not to:
        messages.error(request, 'Tu cuenta no tiene email — no hay a dónde enviar la prueba.')
        return
    connection = None
    from_email = settings.DEFAULT_FROM_EMAIL
    if data['enabled'] and data['host']:
        connection = get_connection(
            'django.core.mail.backends.smtp.EmailBackend',
            host=data['host'], port=data['port'], username=data['username'],
            password=data['password'], use_tls=data['use_tls'],
            timeout=getattr(settings, 'EMAIL_TIMEOUT', 10),
        )
        from_email = data['from_email'] or from_email
    try:
        send_mail(
            'CEV: correo de prueba',
            'Si estás leyendo esto, la configuración de correo saliente funciona.',
            from_email, [to], connection=connection, fail_silently=False,
        )
        messages.success(request, f'Correo de prueba enviado a {to}.')
    except Exception as exc:
        messages.error(request, f'No se pudo enviar el correo de prueba: {exc}')


@login_required
def configuracion(request):
    """Configuración de la cuenta, en pestañas: Marca, Datos, Plantillas y Mi cuenta.

    Marca/Datos/Plantillas son de la EMPRESA cuando la cuenta pertenece a una (las comparte
    todo el equipo, y solo el Administrador las cambia — un usuario cualquiera no debe
    poder cambiarle el logo a los documentos de toda la empresa); en una cuenta personal son
    suyas. «Mi cuenta» siempre es personal: cada uno edita su nombre, tenga empresa o no.
    """
    marca = marca_de(request.user, crear=puede_editar_marca(request.user))
    editable = puede_editar_marca(request.user)

    # La pestaña activa vuelve tras guardar, para no perder de vista dónde estabas.
    tab = request.POST.get('tab') or request.GET.get('tab') or ('marca' if editable else 'plantillas')
    # Marca y Datos solo existen para el Administrador; un ?tab=marca ajeno cae en Plantillas.
    if not editable and tab in ('marca', 'datos'):
        tab = 'plantillas'

    form = MarcaForm(instance=marca) if marca else None
    form_perfil = ProfileNameForm(instance=request.user)

    # Plantillas: cualquier usuario define las suyas. El Administrador edita las de la
    # empresa (el default del equipo); un usuario común edita SU fila Marca, cuyas
    # plantillas pisan a las de la empresa al exportar (ver marca.plantilla_de).
    marca_plantillas = marca if editable else marca_propia(request.user)
    # Si la empresa (o la cuenta personal) solo tiene contratado un módulo, no tiene
    # sentido ofrecer la plantilla del otro.
    modulos = modulos_activos(request.user)
    form_plantillas = PlantillasForm(instance=marca_plantillas, modulos_activos=modulos)

    if request.method == 'POST':
        accion = request.POST.get('action')

        if accion == 'perfil':
            form_perfil = ProfileNameForm(request.POST, instance=request.user)
            if form_perfil.is_valid():
                form_perfil.save()
                messages.success(request, 'Tus datos fueron actualizados.')
                return redirect(f"{reverse('accounts:configuracion')}?tab=cuenta")
            messages.error(request, 'Revisa tus datos.')
        elif tab == 'plantillas':
            if accion in ('vista_previa_proyectos', 'vista_previa_servicios'):
                modulo = 'proyectos' if accion == 'vista_previa_proyectos' else 'servicios'
                return _vista_previa_marca(request, marca, modulo)
            instancia = marca_plantillas if editable else marca_propia(request.user, crear=True)
            form_plantillas = PlantillasForm(
                request.POST, request.FILES, instance=instancia, modulos_activos=modulos,
            )
            if form_plantillas.is_valid():
                form_plantillas.save()
                messages.success(request, 'Plantillas actualizadas.')
                return redirect(f"{reverse('accounts:configuracion')}?tab=plantillas")
            messages.error(request, 'Revisa los datos.')
        else:
            # Todo lo demás toca la marca de la empresa: exige permiso de edición.
            if not editable:
                raise PermissionDenied
            if accion == 'vista_previa':
                return _vista_previa_marca(request, marca)

            form = MarcaForm(request.POST, request.FILES, instance=marca)
            if form.is_valid():
                form.save()
                messages.success(request, 'Configuración actualizada.')
                return redirect(f"{reverse('accounts:configuracion')}?tab={tab}")
            messages.error(request, 'Revisa los datos.')

    return render(request, 'accounts/configuracion.html', {
        'form': form,
        'form_perfil': form_perfil,
        'form_plantillas': form_plantillas,
        'marca': marca,
        'marca_plantillas': marca_plantillas,
        'editable': editable,
        'empresa': get_user_empresa(request.user),
        'tab': tab,
    })


def _vista_previa_marca(request, marca, modulo='proyectos'):
    """Genera un .docx de muestra con la marca aplicada, con la plantilla de `modulo`
    ('proyectos' o 'servicios') — cada módulo tiene la suya, y son independientes.

    Es lo que hace que la pantalla no sea decorativa: se ve el resultado real (logo,
    colores, datos) sin tener que crear un proyecto o servicio y exportarlo para
    comprobarlo. Usa los mismos helpers y la misma precedencia que la exportación real
    (`proyectos.views.exportar_proyecto_word_view` / `servicios.views.exportar_servicio_word_view`).
    """
    import os

    from django.conf import settings as dj_settings
    from django.http import HttpResponse
    from docx import Document

    if modulo == 'servicios':
        nombre_app, nombre_archivo = 'servicios', 'template_servicios.docx'
    else:
        modulo, nombre_app, nombre_archivo = 'proyectos', 'proyectos', 'template_especificaciones.docx'

    plantilla = os.path.join(
        dj_settings.BASE_DIR, nombre_app, 'templates', 'word_templates', nombre_archivo,
    )
    # Misma precedencia que la exportación real: plantilla y logo del usuario → empresa → sistema.
    fuente = plantilla_de(request.user, modulo, plantilla)
    fuente = con_logo(fuente, logo_de(request.user))

    try:
        doc = Document(fuente)
    except Exception:
        messages.error(request, 'No se pudo abrir la plantilla. Revisa que el .docx sea válido.')
        return redirect('accounts:configuracion')

    # Placeholders de ejemplo: se incluyen los de las dos apps a la vez, así una misma
    # plantilla de muestra sirve para cualquiera de las dos — cada .docx solo tiene los
    # marcadores que le corresponden, el resto simplemente no aparece.
    reemplazos = {
        '<<PROYECTO>>': 'PROYECTO DE EJEMPLO',
        '<<TITULO>>': 'SERVICIO DE EJEMPLO',
        '<<SOLICITANTE>>': 'Solicitante de ejemplo',
        '<<CATEGORIA>>': 'Categoría de ejemplo',
        '<<SUBCATEGORIA>>': 'Subcategoría de ejemplo',
        '<<CODIGO>>': 'COD-001',
        '<<UBICACION>>': 'Ubicación de ejemplo',
        '<<USUARIO>>': request.user.get_full_name() or request.user.username,
        '<<TIPO_CONTRATO>>': 'Contrato de Obra (CO)',
        '<<FECHA>>': timezone.now().strftime('%d/%m/%Y'),
        '<<REV>>': '1',
        '<<REV.>>': '1',
        **placeholders_de(marca),
    }
    reemplazar_placeholders(doc, reemplazos)
    insertar_logo(doc, logo_de(request.user))
    aplicar_colores(doc, marca)

    doc.add_heading('Así se ven los títulos de nivel 1', level=1)
    doc.add_heading('Así se ven los de nivel 2', level=2)
    doc.add_heading('Y los subtítulos, en el color secundario', level=3)
    doc.add_paragraph(
        'Este documento es solo una muestra de cómo quedan el logo, los colores y los datos '
        'que configuraste. No corresponde a ningún proyecto ni servicio real.'
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    respuesta = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="vista_previa_{modulo}.docx"'
    return respuesta


@_superuser_required
def email_config(request):
    """Config de correo editable solo por el superuser: SMTP de BD (pisa al .env si
    `enabled`)."""
    config = EmailConfig.load()
    form = EmailConfigForm(instance=config)

    if request.method == 'POST':
        form = EmailConfigForm(request.POST, instance=config)
        if request.POST.get('action') == 'test':
            if form.is_valid():
                _send_test_email(request, form.cleaned_data)
            else:
                messages.error(request, 'Revisa los datos antes de probar el envío.')
        elif form.is_valid():
            form.save()
            messages.success(request, 'Configuración de correo actualizada.')
            return redirect('accounts:email_config')
        else:
            messages.error(request, 'Revisa los datos.')

    return render(request, 'accounts/email_config.html', {'form': form, 'config': config})


@login_required
@require_POST
def impersonate(request):
    """Impersonar a un usuario real. SOLO superuser, en dev y también en producción.

    Guarda su id en la sesión; `ImpersonationMiddleware` reemplaza `request.user` por ese
    usuario en cada request siguiente, así se ve la app exactamente como la ve él (sus
    proyectos, su rol, sus créditos). Sirve para reproducir el problema que reporta un
    cliente sin pedirle la contraseña.

    A quien no es superuser se le responde 404, no 403: un 403 confirmaría que la
    funcionalidad existe.

    Cada entrada y salida queda registrada en el log (nivel WARNING): impersonar en
    producción es potente y no debe poder hacerse sin dejar rastro de quién actuó como quién.

    `request.real_user` es el superuser real cuando ya se está impersonando a alguien (lo
    cuelga el middleware); se usa acá en vez de `request.user` para permitir cambiar de
    usuario impersonado sin tener que salir primero."""
    real_user = getattr(request, 'real_user', request.user)
    if not real_user.is_authenticated or not real_user.is_superuser:
        raise Http404

    user_id = request.POST.get('user_id', '')
    target = User.objects.filter(pk=user_id, is_active=True).first() if user_id else None
    if target:
        # Impersonar a otro superuser no aporta nada y enturbia la auditoría.
        if target.is_superuser:
            messages.error(request, 'No se puede impersonar a otro superuser.')
        else:
            request.session['impersonate_id'] = str(target.pk)
            logger.warning(
                'IMPERSONACIÓN INICIADA: %s (id=%s) actúa como %s (id=%s)',
                real_user.email or real_user.username, real_user.pk,
                target.email or target.username, target.pk,
            )
    else:
        anterior = request.session.pop('impersonate_id', None)
        if anterior:
            logger.warning(
                'IMPERSONACIÓN TERMINADA: %s (id=%s) dejó de actuar como id=%s',
                real_user.email or real_user.username, real_user.pk, anterior,
            )

    next_url = request.POST.get('next', '')
    if not url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()}):
        next_url = ''

    # A quién le va a tocar abrir esa página: al impersonado si acabamos de entrar, al
    # superuser real si acabamos de salir.
    quien = target or real_user

    # El home es el único destino que siempre funciona para cualquiera. Se usa cuando la
    # página de la que venimos no existe para esa cuenta: impersonar desde Cuentas (que es
    # superuser-only) o desde un módulo que la otra persona no tiene habilitado la recibiría
    # con un 403 en la cara. Ojo: /proyectos/ NO sirve de fallback — también da 403 a quien
    # no tenga ese módulo.
    if not next_url or not puede_acceder_a(quien, next_url):
        next_url = reverse('home')

    return redirect(next_url)
